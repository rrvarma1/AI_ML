import asyncio,io,json,os,re,uuid,traceback,html
from datetime import datetime,timedelta,timezone
from pathlib import Path
from typing import Any
import httpx
from bs4 import BeautifulSoup
from urllib.parse import quote_plus, urljoin, urlparse
from docx import Document
from dotenv import load_dotenv, dotenv_values
from fastapi import FastAPI,File,HTTPException,UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openai import OpenAI
from pydantic import BaseModel
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parent.parent
# Load secrets/configuration from the project-root .env file.
ENV_FILE=ROOT/'.env'
load_dotenv(ENV_FILE, override=True)

def get_openai_api_key():
 """Resolve the API key at request time so .env changes are picked up reliably.

 The key is never logged or returned. Reading the file explicitly also avoids
 Windows working-directory/reloader issues.
 """
 # Re-read the project-root .env on every tailoring request.
 if ENV_FILE.exists():
  try:
   values=dotenv_values(ENV_FILE)
   key=(values.get('OPENAI_API_KEY') or '').strip()
   if key:return key
  except Exception:
   pass
 # Fall back to the process environment.
 return (os.environ.get('OPENAI_API_KEY') or '').strip()

def openai_config_error():
 candidates=[]
 if (ROOT/'.env.txt').exists():candidates.append(str(ROOT/'.env.txt'))
 if (ROOT/'env').exists():candidates.append(str(ROOT/'env'))
 hint=f" Checked: {ENV_FILE}."
 if candidates:hint+=f" Found possible misnamed file: {', '.join(candidates)}. Rename it to .env."
 return 'OPENAI_API_KEY was not found.'+hint+' Ensure the file contains a line exactly like OPENAI_API_KEY=sk-... and restart Career Compass.'
DATA=ROOT/'data'; RES=DATA/'resumes'; OUT=DATA/'generated'; STATE=DATA/'portal.json'; PROMPT=ROOT/'prompts'/'resume_rewrite_prompt.txt'
for x in (DATA,RES,OUT): x.mkdir(parents=True,exist_ok=True)
DEFAULT={'organizations':[],'roles':['Software Engineer','Data Analyst'],'jobs':[],'submissions':[],'resume':None,'google_seen':{},'google_baselined':False}
app=FastAPI(title='Career Compass')
def state():
 if not STATE.exists(): save(DEFAULT.copy())
 d=json.loads(STATE.read_text()); [d.setdefault(k,v) for k,v in DEFAULT.items()]; return d
def save(d):
 t=STATE.with_suffix('.tmp');t.write_text(json.dumps(d,indent=2));t.replace(STATE)
def text(x): return BeautifulSoup(x or '','html.parser').get_text(' ',strip=True)
def clean_description(x):
 """Convert ATS HTML/escaped HTML into compact plain text for dashboard summaries."""
 s=str(x or '')
 # Some ATS feeds (notably Greenhouse) return HTML entities containing HTML markup.
 # Decode a few times until stable, then strip all tags.
 for _ in range(3):
  decoded=html.unescape(s)
  if decoded==s: break
  s=decoded
 s=BeautifulSoup(s,'html.parser').get_text(' ',strip=True)
 s=html.unescape(s)
 s=re.sub(r'\s+',' ',s).strip()
 # Remove a leading standalone requisition marker such as (P-1384).
 s=re.sub(r'^\([A-Za-z]-?\d{3,}\)\s*','',s)
 return s


def summarize_description(x, max_lines=3):
 """Create a compact 3-line extractive summary from a cleaned job description."""
 s=clean_description(x)
 if not s:return ''
 # Split into sentence-like units, then prefer informative lines over boilerplate/headings.
 parts=[re.sub(r'\s+',' ',q).strip(' -•\t') for q in re.split(r'(?<=[.!?])\s+|\s+[•|]\s+',s) if q.strip()]
 boiler=re.compile(r'^(about (the )?(team|role|company)|who we are|what we do|equal opportunity|benefits|compensation|pay range)\b',re.I)
 keywords=('responsib','lead','build','manage','drive','develop','deliver','partner','design','architect','experience','skills','qualif','customer','platform','strategy','data','engineering','program')
 scored=[]
 for i,q in enumerate(parts):
  if len(q)<30:continue
  score=sum(1 for k in keywords if k in q.lower())
  if boiler.search(q):score-=2
  score+=max(0,3-i)*0.05
  scored.append((score,i,q))
 chosen=[]
 for _,i,q in sorted(scored,key=lambda t:(-t[0],t[1])):
  # Avoid near-duplicate sentences.
  qwords=set(re.findall(r'[a-z]{4,}',q.lower()))
  if any(len(qwords & w)/max(1,len(qwords|w))>.55 for _,w in chosen):continue
  chosen.append((i,qwords));
  if len(chosen)>=max_lines:break
 selected=[parts[i] for i,_ in sorted(chosen,key=lambda t:t[0])]
 if not selected:selected=parts[:max_lines]
 lines=[]
 for q in selected[:max_lines]:
  q=q.strip()
  if len(q)>155:q=q[:152].rsplit(' ',1)[0]+'…'
  lines.append(q)
 return '\n'.join(lines)

def date(x):
 try:return datetime.fromisoformat(str(x).replace('Z','+00:00')).replace(tzinfo=timezone.utc)
 except:return None
def recent(x):
 d=date(x);return bool(d and timedelta()<=datetime.now(timezone.utc)-d<=timedelta(days=7))
def words(x):return set(re.findall(r'[a-z0-9+#.]{2,}',x.lower()))-{'and','the','for','with','senior','junior','lead'}
F=[{'software','engineer','developer','programmer'},{'data','analyst','analytics','bi'},{'product','manager','pm'},{'designer','design','ux','ui'},{'marketing','growth','brand'},{'finance','financial','accounting'},{'recruiter','recruiting','talent'}]
def match(title,roles):
 a=words(title)
 for r in roles:
  b=words(r)
  if ' '.join(r.lower().split()) in ' '.join(title.lower().split()) or len(a&b)>=max(1,min(2,len(b))) or any(a&f and b&f for f in F):return True
 return False

def norm_location(x):
 return re.sub(r'[^a-z0-9]+',' ',(x or '').lower()).strip()
LOCATION_ALIASES={
 'bengaluru':{'bengaluru','bangalore'},
 'bangalore':{'bengaluru','bangalore'},
 'new york':{'new york','new york city','nyc'},
 'san francisco':{'san francisco','sf'},
}
def location_terms(preferred, normalized=''):
 vals={norm_location(preferred),norm_location(normalized)}-{''}
 expanded=set(vals)
 for v in list(vals): expanded |= LOCATION_ALIASES.get(v,set())
 return {x for x in expanded if x}
def location_match(job_location, preferred, normalized=''):
 """Strict city filter. A job must explicitly contain the preferred city/alias.

 Do not reverse-match the whole job location against a short preferred value;
 that previously allowed over-broad card text to leak unrelated locations.
 """
 jl=norm_location(job_location)
 if not jl:return False
 for city in location_terms(preferred,normalized):
  if re.search(r'(?<![a-z0-9])'+re.escape(city)+r'(?![a-z0-9])',jl):return True
 return False

async def validate_city(value):
 """Validate and normalize a user-entered city with OpenStreetMap Nominatim."""
 q=(value or '').strip()
 if not q:return None
 try:
  async with httpx.AsyncClient(timeout=12,follow_redirects=True,headers={'User-Agent':'CareerCompass/1.0 (local job preference validator)'}) as c:
   r=await c.get('https://nominatim.openstreetmap.org/search',params={'q':q,'format':'jsonv2','limit':5,'addressdetails':1})
   r.raise_for_status(); results=r.json()
 except Exception as e:
  raise HTTPException(503,f"Could not validate location '{q}' right now. Check your internet connection and try again.") from e
 for row in results:
  a=row.get('address') or {}
  city=a.get('city') or a.get('town') or a.get('municipality') or a.get('village') or a.get('borough')
  typ=(row.get('type') or '').lower()
  if city or typ in {'city','town','municipality','village','borough'}:
   canonical=city or row.get('name') or q
   return {'input':q,'city':canonical,'display_name':row.get('display_name') or canonical}
 return None

KNOWN_ATS_BOARDS={
 'databricks':('greenhouse','databricks'),
}

def _token_from_url(url, source):
 u=(url or '').strip()
 if not u:return ''
 patterns={
  'greenhouse':[
   r'https?://(?:boards|job-boards)\.greenhouse\.io/([^/?#]+)',
   r'https?://boards-api\.greenhouse\.io/v1/boards/([^/?#]+)',
  ],
  'lever':[
   r'https?://jobs\.lever\.co/([^/?#]+)',
   r'https?://api\.lever\.co/v0/postings/([^/?#]+)',
  ],
 }
 for pat in patterns.get(source,[]):
  m=re.search(pat,u,re.I)
  if m:return m.group(1)
 return ''

def _local_source_detection(name,url):
 """Detect ATS from organization name/URL without network access."""
 n=(name or '').strip().lower();u=(url or '').strip();ul=u.lower()
 for key,(source,token) in KNOWN_ATS_BOARDS.items():
  if key in n or key in ul:return {'source_type':source,'board_token':token,'detected_by':'known organization'}
 if 'greenhouse.io' in ul:
  return {'source_type':'greenhouse','board_token':_token_from_url(u,'greenhouse'),'detected_by':'careers URL'}
 if 'lever.co' in ul:
  return {'source_type':'lever','board_token':_token_from_url(u,'lever'),'detected_by':'careers URL'}
 if 'google.com/about/careers' in ul:
  return {'source_type':'google_careers','board_token':'','detected_by':'careers URL'}
 return {'source_type':'generic','board_token':'','detected_by':'default'}

async def detect_source(name,url):
 """Detect supported ATS and board token, including company-hosted careers pages."""
 base=_local_source_detection(name,url)
 if base['source_type']!='generic' or not (url or '').strip():return base
 try:
  async with httpx.AsyncClient(timeout=12,follow_redirects=True,headers={'User-Agent':'Mozilla/5.0 CareerCompass/1.0'}) as c:
   r=await c.get(url)
   final=str(r.url);html=r.text[:1500000]
  # Redirects sometimes reveal the ATS directly.
  redirected=_local_source_detection(name,final)
  if redirected['source_type']!='generic':
   redirected['detected_by']='redirected careers URL';return redirected
  # Company-hosted pages often embed links to the underlying ATS.
  gh=re.search(r'https?://(?:boards|job-boards)\.greenhouse\.io/([^/?#"\\\'<> ]+)',html,re.I)
  if gh:return {'source_type':'greenhouse','board_token':gh.group(1),'detected_by':'careers page'}
  lv=re.search(r'https?://jobs\.lever\.co/([^/?#"\\\'<> ]+)',html,re.I)
  if lv:return {'source_type':'lever','board_token':lv.group(1),'detected_by':'careers page'}
  # Greenhouse-hosted application widgets frequently expose gh_jid even on custom domains.
  if re.search(r'\bgh_jid\b',html,re.I):
   known=_local_source_detection(name,url)
   if known.get('board_token'):return known
 except Exception as e:
  print(f"[SOURCE] Detection fallback for {name}: {type(e).__name__}: {e}")
 return base

def record(org,raw,source):
 title=raw.get('title') or raw.get('name') or ''; posted=raw.get('updated_at') or raw.get('createdAt') or raw.get('datePosted') or raw.get('published_at')
 if not title or not recent(posted):return None
 desc=raw.get('content') or raw.get('description') or raw.get('descriptionPlain') or ''; url=raw.get('absolute_url') or raw.get('hostedUrl') or raw.get('url') or org.get('careers_url',''); ident=raw.get('id') or raw.get('requisitionId') or raw.get('jobId') or url
 loc=raw.get('location') or raw.get('locationName') or (raw.get('categories') or {}).get('location') or raw.get('workplaceType') or ''; loc=loc.get('name','') if isinstance(loc,dict) else loc; return {'id':f'{source}:{ident}','organization':org['name'],'logo_url':org.get('logo_url',''),'job_id':str(ident),'title':text(title),'description':clean_description(desc),'summary':summarize_description(desc),'location':text(str(loc)),'source_url':url,'posted_at':date(posted).isoformat(),'source':source}
async def google_httpx_links(c, url):
 """Try Google's server-rendered results first; browser-like headers improve SSR responses."""
 r=await c.get(url,headers={
  'User-Agent':'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36',
  'Accept':'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8',
  'Accept-Language':'en-US,en;q=0.9'
 })
 r.raise_for_status()
 soup=BeautifulSoup(r.text,'html.parser')
 out=[]
 for a in soup.select('a[href]'):
  href=urljoin(str(r.url),a.get('href',''))
  if re.search(r'/about/careers/applications/jobs/results/\d+[-/]',href):
   title=(a.select_one('h3') or a.select_one('h2'))
   label=text(title.get_text(' ',strip=True) if title else a.get_text(' ',strip=True))
   if label: out.append((href,label,text(a.parent.get_text(' ',strip=True) if a.parent else label),''))
 return r.status_code,out

def _google_playwright_links_subprocess(url):
 """Run Playwright in an isolated Python process.

 Uvicorn on Windows may select an asyncio event loop that cannot create
 subprocesses. Playwright itself needs subprocess support even when using its
 sync API, so the browser worker is launched in a separate process where we
 explicitly select WindowsProactorEventLoopPolicy before Playwright starts.
 """
 import subprocess
 import sys
 worker=ROOT/'app'/'google_browser_worker.py'
 cmd=[sys.executable,str(worker),url]
 proc=subprocess.run(
  cmd,
  capture_output=True,
  text=True,
  timeout=75,
  creationflags=(subprocess.CREATE_NO_WINDOW if sys.platform=='win32' and hasattr(subprocess,'CREATE_NO_WINDOW') else 0),
 )
 stdout=(proc.stdout or '').strip()
 stderr=(proc.stderr or '').strip()
 if stderr:
  print(f"[GOOGLE][PLAYWRIGHT][WORKER-STDERR] {stderr[-2000:]}")
 if not stdout:
  raise RuntimeError(f'Google browser worker returned no output (exit={proc.returncode}).')
 # The worker writes one JSON object. Be defensive if a dependency prints noise.
 line=stdout.splitlines()[-1]
 try:
  payload=json.loads(line)
 except Exception as e:
  raise RuntimeError(f'Unable to parse Google browser worker output: {stdout[-2000:]}') from e
 if not payload.get('ok'):
  raise RuntimeError(payload.get('error') or f'Google browser worker failed (exit={proc.returncode}).')
 out=[]
 for row in payload.get('rows',[]):
  href=row.get('href','')
  label=text(row.get('text',''))
  lines=[x.strip() for x in label.splitlines() if x.strip()]
  title=lines[0] if lines else label
  if title:
   out.append((href,title,text(row.get('parent') or label),text(row.get('location') or '')))
 return out

async def google_playwright_links(url):
 """Launch the isolated browser worker without blocking FastAPI's event loop."""
 try:
  return await asyncio.to_thread(_google_playwright_links_subprocess,url)
 except Exception as e:
  print(f"[GOOGLE][PLAYWRIGHT][ERROR] {type(e).__name__}: {e}")
  raise

async def fetch_google(org, roles):
 base=(org.get('careers_url') or 'https://www.google.com/about/careers/applications/jobs/results').rstrip('/')
 raw=[]; seen=set(); diagnostics=[]
 async with httpx.AsyncClient(timeout=30,follow_redirects=True) as c:
  for role in roles or ['']:
   for page_no in (1,2):
    location=(org.get('location_normalized') or org.get('location') or '').strip()
    url=f"{base}?q={quote_plus(role)}&sort_by=date"
    if location:url += f"&location={quote_plus(location)}"
    if page_no>1:url += f"&page={page_no}"
    print(f"[GOOGLE] Fetch role={role!r} page={page_no} | {url}")
    status,links=await google_httpx_links(c,url)
    mode='httpx'
    print(f"[GOOGLE] HTTP {status}; static parser found {len(links)} job links")
    if not links:
     print('[GOOGLE] Static HTML has no job links; rendering page with Playwright/Chromium...')
     links=await google_playwright_links(url);mode='playwright'
     print(f"[GOOGLE] Browser-rendered parser found {len(links)} job links")
    added=0
    for href,title,desc,location in links:
     m=re.search(r'/jobs/results/(\d+)(?:-([^?#/]+))?',href)
     ident=m.group(1) if m else href
     if ident in seen: continue
     seen.add(ident);added+=1
     raw.append({'id':ident,'title':title,'description':desc,'location':location,'url':href})
    diagnostics.append({'role':role,'page':page_no,'status':status,'mode':mode,'links':len(links),'new_unique':added})
    if page_no==1 and not links: break
 return raw,diagnostics

async def feed(org,roles=None):
 typ=org['source_type']
 # Databricks' public careers site is backed by Greenhouse (job links use gh_jid).
 # Users can therefore paste the normal Databricks careers URL and leave Source as
 # "Other careers site"; Career Compass will transparently use the Greenhouse feed.
 is_databricks=('databricks' in (org.get('name') or '').lower() or 'databricks.com/company/careers' in (org.get('careers_url') or '').lower())
 if typ=='generic' and is_databricks:
  typ='greenhouse'; org=dict(org); org['board_token']='databricks'
  print('[SOURCE] Auto-detected Databricks Greenhouse board token=databricks')
 if typ in ('google_careers','google'):
  raw,diagnostics=await fetch_google(org,roles or [])
  now=datetime.now(timezone.utc).isoformat()
  jobs=[]
  for x in raw:
   jobs.append({'id':'google_careers:'+str(x['id']),'organization':org['name'],'logo_url':org.get('logo_url',''),'job_id':str(x['id']),'title':text(x['title']),'description':clean_description(x['description']),'location':text(x.get('location','')),'source_url':x['url'],'posted_at':now,'source':'google_careers'})
  return jobs,diagnostics
 async with httpx.AsyncClient(timeout=25,follow_redirects=True,headers={'User-Agent':'Mozilla/5.0 CareerCompass/1.0'}) as c:
  if typ=='greenhouse': raw=(await c.get(f"https://boards-api.greenhouse.io/v1/boards/{org['board_token']}/jobs?content=true")).json().get('jobs',[])
  elif typ=='lever':raw=(await c.get(f"https://api.lever.co/v0/postings/{org['board_token']}?mode=json")).json()
  else:
   html=(await c.get(org['careers_url'])).text; raw=[]
   for e in BeautifulSoup(html,'html.parser').select('script[type="application/ld+json"]'):
    try:
     d=json.loads(e.string or ''); raw += d if isinstance(d,list) else d.get('@graph',[d])
    except:pass
 return [record(org,x,typ) for x in raw if isinstance(x,dict) and (typ!='generic' or x.get('@type')=='JobPosting' or 'JobPosting' in x.get('@type',[]))],[]
class Config(BaseModel): organizations:list[dict[str,Any]]=[];roles:list[str]=[]
class IDs(BaseModel): job_ids:list[str]
class P(BaseModel): text:str
@app.get('/api/config')
def config():
 d=state();return {'organizations':d['organizations'],'roles':d['roles'],'resume':d['resume']}
@app.post('/api/source-detect')
async def source_detect(payload:dict[str,Any]):
 name=(payload.get('name') or '').strip();url=(payload.get('careers_url') or '').strip()
 if not url:return {'source_type':'generic','board_token':'','detected_by':'default'}
 result=await detect_source(name,url)
 return result

@app.put('/api/config')
async def putconfig(p:Config):
 validated=[]
 for o in p.organizations:
  if not o.get('name'):raise HTTPException(400,'Every organization needs a name.')
  if not o.get('careers_url'):raise HTTPException(400,f"Enter a careers URL for {o.get('name','this organization')}.")
  detected=await detect_source(o.get('name',''),o.get('careers_url',''))
  o=dict(o);o['source_type']=detected['source_type'];o['board_token']=detected.get('board_token','')
  o['source_detected_by']=detected.get('detected_by','automatic detection')
  if o['source_type'] in ('greenhouse','lever') and not o.get('board_token'):
   raise HTTPException(400,f"Career Compass detected {o['source_type'].title()} for {o['name']} but could not determine its board token automatically. Check the careers URL and try again.")
  location=(o.get('location') or '').strip()
  if not location:raise HTTPException(400,f"Enter a location for {o.get('name','this organization')}.")
  city=await validate_city(location)
  if not city:raise HTTPException(400,f"'{location}' does not appear to be a valid city. Enter a city such as Bengaluru, London, or New York.")
  clean=dict(o);clean['location']=location;clean['location_normalized']=city['city'];clean['location_display']=city['display_name'];validated.append(clean)
 d=state();d['organizations']=validated;d['roles']=[x.strip() for x in p.roles if x.strip()];save(d);return {'ok':True,'organizations':validated}
@app.get('/api/jobs')
def getjobs():return {'jobs':state()['jobs']}
@app.post('/api/refresh')
async def refresh():
 d=state();found=[];errors=[];diagnostics=[]
 now=datetime.now(timezone.utc).isoformat()
 for o in d['organizations']:
  print(f"\n[REFRESH] Organization={o.get('name')} source={o.get('source_type')} url={o.get('careers_url','')}")
  try:
   raw,source_diag=await feed(o,d['roles'])
   print(f"[REFRESH] {o['name']}: raw={len(raw)}")
   matched=[x for x in raw if x and match(x['title'],d['roles'])]
   print(f"[REFRESH] {o['name']}: role-matched={len(matched)}")
   preferred=(o.get('location') or '').strip(); normalized=(o.get('location_normalized') or '').strip()
   if preferred:
    location_matched=[]
    for x in matched:
     original_location=(x.get('location') or '').strip()
     # Google has already been queried with the preferred location. If its
     # rendered card omits location text, accept the result using the Google
     # location filter; otherwise require an explicit strict city match.
     google_source=o['source_type'] in ('google_careers','google')
     if (google_source and not original_location) or location_match(original_location,preferred,normalized):
      x['source_location']=original_location
      x['location']=normalized or preferred
      location_matched.append(x)
   else:
    location_matched=matched
   print(f"[REFRESH] {o['name']}: location={preferred!r} matched={len(location_matched)}")
   matched=location_matched
   if o['source_type'] in ('google_careers','google'):
    seen=d.setdefault('google_seen',{})
    for x in matched:
     if x['id'] not in seen: seen[x['id']]=now
     x['posted_at']=seen[x['id']]
    cutoff=datetime.now(timezone.utc)-timedelta(days=7)
    recent_matches=[x for x in matched if (date(seen.get(x['id'])) or datetime.now(timezone.utc))>=cutoff]
    found += recent_matches
    diagnostics.append({'organization':o['name'],'source':'google_careers','raw':len(raw),'matched':len(matched),'preferred_location':preferred,'kept_last_7_days':len(recent_matches),'fetches':source_diag})
    print(f"[REFRESH] {o['name']}: kept first-seen<=7d={len(recent_matches)}")
   else:
    found += matched
    diagnostics.append({'organization':o['name'],'source':o['source_type'],'raw':len(raw),'matched':len(matched),'preferred_location':preferred,'fetches':source_diag})
  except Exception as e:
   msg=f"{o.get('name','Unknown')}: {e}"
   errors.append(msg);diagnostics.append({'organization':o.get('name'),'source':o.get('source_type'),'error':str(e)})
   print(f"[REFRESH][ERROR] {msg}")
   traceback.print_exc()
 d['jobs']=sorted({x['id']:x for x in found}.values(),key=lambda x:x['posted_at'],reverse=True)
 save(d)
 print(f"[REFRESH] FINAL jobs stored={len(d['jobs'])} errors={len(errors)}")
 return {'jobs':d['jobs'],'errors':errors,'checked':len(d['organizations']),'diagnostics':diagnostics}
@app.post('/api/submissions')
def submit(p:IDs):
 d=state();existing={x['job']['id'] for x in d['submissions']};chosen=[x for x in d['jobs'] if x['id'] in p.job_ids]; added=0
 for x in chosen:
  if x['id'] not in existing:d['submissions'].insert(0,{'id':str(uuid.uuid4()),'job':x,'submitted_at':datetime.now(timezone.utc).isoformat(),'tailored_files':[]});added+=1
 save(d);return {'added':added}
@app.get('/api/submissions')
def history():return {'submissions':state()['submissions']}
def resume_text(name,b):
 s=Path(name).suffix.lower()
 if s=='.txt':return b.decode(errors='ignore')
 if s=='.pdf':return '\n'.join(x.extract_text() or '' for x in PdfReader(io.BytesIO(b)).pages)
 if s=='.docx':return '\n'.join(x.text for x in Document(io.BytesIO(b)).paragraphs)
 raise HTTPException(400,'Upload TXT, PDF, or DOCX.')
@app.post('/api/resume')
async def upload(file:UploadFile=File(...)):
 b=await file.read();v=resume_text(file.filename,b).strip()
 if not v:raise HTTPException(400,'No readable text was found.')
 safe=re.sub(r'[^A-Za-z0-9._-]','_',file.filename);dest=RES/f'{uuid.uuid4().hex}_{safe}';dest.write_bytes(b);d=state();d['resume']={'filename':file.filename,'path':dest.name,'text':v};save(d);return {'filename':file.filename,'characters':len(v)}
@app.get('/api/prompt')
def getprompt():return {'text':PROMPT.read_text()}
@app.put('/api/prompt')
def putprompt(p:P):
 if '{role_name}' not in p.text or '{organization_name}' not in p.text:raise HTTPException(400,'Prompt must include {role_name} and {organization_name}.')
 PROMPT.write_text(p.text.strip()+'\n');return {'ok':True}
def safe(x):return re.sub(r'[^A-Za-z0-9]+','_',x).strip('_')[:45] or 'Role'
def docx(content,path):
 d=Document();d.sections[0].top_margin=d.sections[0].bottom_margin=457200
 for l in content.splitlines():
  l=l.strip()
  if not l:continue
  if l.isupper() or l.rstrip(':') in {'SUMMARY','EXPERIENCE','EDUCATION','SKILLS','CERTIFICATIONS','PROJECTS'}:d.add_heading(l.rstrip(':'),2)
  elif l.startswith(('- ','• ')):d.add_paragraph(l[2:],style='List Bullet')
  else:d.add_paragraph(l)
 d.save(path)
@app.post('/api/tailor')
def tailor(p:IDs):
 d=state();r=d['resume']
 if not r:raise HTTPException(400,'Upload a base resume first.')
 key=get_openai_api_key()
 if not key:raise HTTPException(400,openai_config_error())
 chosen=[x for x in d['jobs'] if x['id'] in p.job_ids]
 if not chosen:raise HTTPException(400,'Select one or more current jobs.')
 client=OpenAI(api_key=key);template=PROMPT.read_text();files=[]
 for j in chosen:
  instruction=template.replace('{role_name}',j['title']).replace('{organization_name}',j['organization']);out=client.responses.create(model=os.getenv('OPENAI_MODEL','gpt-5'),input=f'{instruction}\n\nBASE RESUME:\n{r["text"]}\n\nJOB DESCRIPTION:\n{j["description"]}').output_text
  name=f"Resume_Job_{safe(j['title'])}_{safe(j['organization'])}_{datetime.now().strftime('%Y-%m-%d')}.docx";docx(out,OUT/name);files.append({'job_id':j['id'],'filename':name,'url':'/api/download/'+name})
  for s in d['submissions']:
   if s['job']['id']==j['id']:s['tailored_files'].append(name)
 save(d);return {'files':files}
@app.get('/api/download/{filename}')
def download(filename:str):
 p=OUT/Path(filename).name
 if not p.exists():raise HTTPException(404,'File not found.')
 return FileResponse(p,filename=p.name,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
app.mount('/',StaticFiles(directory=ROOT/'app'/'static',html=True),name='static')
