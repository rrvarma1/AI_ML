from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import sys
import unicodedata
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pymupdf as fitz
from openpyxl import load_workbook

CORPUS = Path(os.environ["RBI_CORPUS_DIR"]).expanduser().resolve()
INVENTORY_FILE = Path(os.environ["RBI_INVENTORY_FILE"]).expanduser().resolve()
OUT = Path(os.environ["RBI_RAW_OUTPUT_DIR"]).expanduser().resolve()
RAW = OUT / "raw_extractions"
ASSETS = OUT / "assets"


def iso_excel_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def clean_key(value):
    return str(value or "").strip().lower().replace(" ", "_")


def normalized_name(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.replace("–", "-").replace("—", "-")
    text = re.sub(r"[\\/:*?\"<>|]", " ", text)
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).lower()
    return " ".join(text.split())


def sha256_file(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def make_serializable(value):
    if isinstance(value, fitz.Rect):
        return [round(x, 3) for x in value]
    if isinstance(value, fitz.Point):
        return [round(value.x, 3), round(value.y, 3)]
    if isinstance(value, bytes):
        return {"byte_length": len(value), "sha256": hashlib.sha256(value).hexdigest()}
    if isinstance(value, dict):
        return {str(k): make_serializable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [make_serializable(v) for v in value]
    return value


def inventory_rows():
    workbook = load_workbook(INVENTORY_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not values:
        raise ValueError(f"Inventory workbook is empty: {INVENTORY_FILE}")
    headers = [clean_key(x) for x in values[0]]
    rows = []
    for raw in values[1:]:
        if not any(x not in (None, "") for x in raw):
            continue
        row = {headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))}
        row = {k: (v.strip() if isinstance(v, str) else v) for k, v in row.items()}
        row["publication_date"] = iso_excel_date(row.get("publication_date"))
        row["updated_through"] = iso_excel_date(row.get("updated_through"))
        row["status"] = str(row.get("status") or "").strip().lower()
        rows.append(row)
    return rows


def match_inventory(pdf, rows):
    key = normalized_name(pdf.name)
    exact = [r for r in rows if normalized_name(r.get("file_name")) == key]
    if len(exact) == 1:
        return exact[0], "normalized_exact"
    pdf_tokens = set(key.split())
    scored = []
    for row in rows:
        candidate = set(normalized_name(row.get("file_name")).split())
        score = len(pdf_tokens & candidate) / max(1, len(pdf_tokens | candidate))
        scored.append((score, row))
    scored.sort(key=lambda x: x[0], reverse=True)
    if scored and scored[0][0] >= 0.82:
        return scored[0][1], f"token_similarity:{scored[0][0]:.3f}"
    return None, "unmatched"


def span_record(span):
    return {
        "text": span.get("text", ""),
        "bbox": [round(x, 3) for x in span.get("bbox", [])],
        "origin": [round(x, 3) for x in span.get("origin", [])],
        "font": span.get("font"),
        "size": round(span.get("size", 0), 3),
        "flags": span.get("flags"),
        "char_flags": span.get("char_flags"),
        "color": span.get("color"),
        "alpha": span.get("alpha"),
        "ascender": round(span.get("ascender", 0), 4),
        "descender": round(span.get("descender", 0), 4),
    }


def block_record(block):
    base = {
        "block_number": block.get("number"),
        "block_type": "text" if block.get("type") == 0 else "image",
        "bbox": [round(x, 3) for x in block.get("bbox", [])],
    }
    if block.get("type") == 0:
        lines = []
        text_parts = []
        for line in block.get("lines", []):
            spans = [span_record(s) for s in line.get("spans", [])]
            line_text = "".join(s["text"] for s in spans)
            text_parts.append(line_text)
            lines.append({
                "bbox": [round(x, 3) for x in line.get("bbox", [])],
                "wmode": line.get("wmode"),
                "direction": [round(x, 4) for x in line.get("dir", [])],
                "text": line_text,
                "spans": spans,
            })
        base.update({"text": "\n".join(text_parts), "lines": lines})
    else:
        image_data = block.get("image", b"")
        base.update({
            "width": block.get("width"),
            "height": block.get("height"),
            "extension": block.get("ext"),
            "colorspace": block.get("colorspace"),
            "xres": block.get("xres"),
            "yres": block.get("yres"),
            "image_byte_length": len(image_data),
            "image_sha256": hashlib.sha256(image_data).hexdigest() if image_data else None,
        })
    return base


def extract_table(page, table, index):
    try:
        data = table.extract()
    except Exception as exc:
        data = None
        error = f"{type(exc).__name__}: {exc}"
    else:
        error = None
    return {
        "table_index": index,
        "bbox": [round(x, 3) for x in table.bbox],
        "row_count": getattr(table, "row_count", None),
        "column_count": getattr(table, "col_count", None),
        "header": make_serializable(getattr(table, "header", None).__dict__) if getattr(table, "header", None) else None,
        "cells": make_serializable(getattr(table, "cells", [])),
        "data": data,
        "extraction_error": error,
    }


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    RAW.mkdir(parents=True)
    ASSETS.mkdir(parents=True)
    rows = inventory_rows()
    pdfs = sorted(CORPUS.glob("*.pdf"))
    manifest = []
    warnings = []

    for pdf in pdfs:
        inv, match_method = match_inventory(pdf, rows)
        document_id = (inv or {}).get("document_id") or re.sub(r"[^a-z0-9]+", "_", pdf.stem.lower()).strip("_")
        doc = fitz.open(pdf)
        doc_assets = ASSETS / document_id
        doc_assets.mkdir(parents=True, exist_ok=True)
        page_records = []
        text_lengths = []
        table_total = image_total = 0
        document_warnings = []
        seen_xrefs = set()

        for page_index, page in enumerate(doc):
            page_no = page_index + 1
            raw = page.get_text("dict", sort=False)
            blocks = [block_record(b) for b in raw.get("blocks", [])]
            plain_text = page.get_text("text", sort=False)
            sorted_text = page.get_text("text", sort=True)
            text_length = len(plain_text.strip())
            text_lengths.append(text_length)
            image_refs = []
            for image in page.get_images(full=True):
                xref = image[0]
                record = {"xref": xref, "smask": image[1], "width": image[2], "height": image[3], "bpc": image[4], "colorspace": image[5], "name": image[7]}
                if xref not in seen_xrefs:
                    try:
                        extracted = doc.extract_image(xref)
                        asset_name = f"xref_{xref}.{extracted['ext']}"
                        asset_path = doc_assets / asset_name
                        asset_path.write_bytes(extracted["image"])
                        record["asset_path"] = str(asset_path.relative_to(OUT))
                        record["sha256"] = hashlib.sha256(extracted["image"]).hexdigest()
                        record["byte_length"] = len(extracted["image"])
                        seen_xrefs.add(xref)
                    except Exception as exc:
                        record["extraction_error"] = f"{type(exc).__name__}: {exc}"
                image_refs.append(record)
            image_total += len(image_refs)

            try:
                found = page.find_tables()
                tables = [extract_table(page, table, i) for i, table in enumerate(found.tables)]
            except Exception as exc:
                tables = []
                document_warnings.append({"page_number": page_no, "type": "table_detection_error", "detail": f"{type(exc).__name__}: {exc}"})
            table_total += len(tables)

            page_warnings = []
            coverage = text_length / max(1.0, page.rect.width * page.rect.height) * 1000
            if text_length == 0:
                page_warnings.append("no_selectable_text")
            elif coverage < 0.15 and image_refs:
                page_warnings.append("possible_scanned_or_image_dominant_page")
            if "\ufffd" in plain_text:
                page_warnings.append("replacement_characters_detected")
            if plain_text and sorted_text and plain_text != sorted_text:
                page_warnings.append("native_and_sorted_reading_order_differ")

            links = []
            for link in page.get_links():
                link = dict(link)
                if "from" in link:
                    link["from"] = make_serializable(link["from"])
                links.append(make_serializable(link))

            page_records.append({
                "page_number": page_no,
                "page_index": page_index,
                "width": round(page.rect.width, 3),
                "height": round(page.rect.height, 3),
                "rotation": page.rotation,
                "plain_text": plain_text,
                "sorted_text": sorted_text,
                "text_character_count": text_length,
                "blocks": blocks,
                "tables": tables,
                "image_references": image_refs,
                "links": links,
                "warnings": page_warnings,
            })
            for warning in page_warnings:
                document_warnings.append({"page_number": page_no, "type": warning})

        metadata = dict(inv or {})
        metadata.update({
            "document_id": document_id,
            "source_file_name": pdf.name,
            "inventory_file_name": (inv or {}).get("file_name"),
            "inventory_match_method": match_method,
            "pdf_sha256": sha256_file(pdf),
            "pdf_byte_length": pdf.stat().st_size,
            "page_count": doc.page_count,
            "pdf_metadata": {k: v for k, v in doc.metadata.items() if v},
            "raw_extraction_format_version": "1.0.0",
            "extracted_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
            "extractor": {"name": "PyMuPDF", "version": fitz.VersionBind},
        })
        record = {"document_metadata": metadata, "pages": page_records}
        raw_path = RAW / f"{document_id}.json"
        raw_path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
        summary = {
            "document_id": document_id,
            "title": metadata.get("title"),
            "source_file_name": pdf.name,
            "raw_file": str(raw_path.relative_to(OUT)),
            "inventory_match_method": match_method,
            "page_count": doc.page_count,
            "pages_with_no_text": sum(x == 0 for x in text_lengths),
            "pages_with_low_text": sum(0 < x < 50 for x in text_lengths),
            "text_character_count": sum(text_lengths),
            "table_count": table_total,
            "image_reference_count": image_total,
            "warning_count": len(document_warnings),
            "raw_json_sha256": sha256_file(raw_path),
        }
        manifest.append(summary)
        warnings.extend({"document_id": document_id, **w} for w in document_warnings)
        doc.close()

    matched_ids = {m["document_id"] for m in manifest}
    expected_ids = {r["document_id"] for r in rows}
    validation = {
        "generated_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "expected_inventory_documents": len(rows),
        "pdf_files_found": len(pdfs),
        "raw_files_generated": len(manifest),
        "matched_inventory_documents": len(matched_ids & expected_ids),
        "missing_inventory_document_ids": sorted(expected_ids - matched_ids),
        "unexpected_document_ids": sorted(matched_ids - expected_ids),
        "total_pages": sum(m["page_count"] for m in manifest),
        "total_characters": sum(m["text_character_count"] for m in manifest),
        "total_tables": sum(m["table_count"] for m in manifest),
        "total_image_references": sum(m["image_reference_count"] for m in manifest),
        "total_warnings": len(warnings),
        "checks": {
            "ten_inventory_rows": len(rows) == 10,
            "ten_pdfs": len(pdfs) == 10,
            "ten_raw_files": len(manifest) == 10,
            "all_inventory_rows_matched": expected_ids == matched_ids,
            "all_raw_files_nonempty": all((OUT / m["raw_file"]).stat().st_size > 100 for m in manifest),
        },
        "documents": manifest,
    }
    (OUT / "document_manifest.json").write_text(json.dumps({"documents": manifest}, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "extraction_warnings.jsonl").write_text("".join(json.dumps(w, ensure_ascii=False) + "\n" for w in warnings), encoding="utf-8")
    (OUT / "validation_report.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "README.md").write_text(
        "# RBI Raw Extraction Layer\n\n"
        "Loss-minimized page/block extraction of the supplied RBI PDF corpus. Each document JSON preserves document metadata, page dimensions, native and sorted text, text blocks, lines, font spans, bounding boxes, table detections, image references, links, and page-level warnings. Extracted embedded images are stored under `assets/`. No semantic chunking or text cleanup has been applied.\n",
        encoding="utf-8",
    )
    print(json.dumps(validation, indent=2))


if __name__ == "__main__":
    main()
